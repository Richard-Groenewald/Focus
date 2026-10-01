"""Parse finance's invoice workbooks into billing schedules (Focus v7.9.89, step 1 of billing).

Richard 2026-10-01: finance builds every contract invoice as one Excel sheet (master file + the
Harmony workbook), types a summary of each into a Pastel Partner import file, and imports it. This
reads every invoice sheet into a schedule (header fields + itemised lines), joins the Pastel fields
(GL code, cost centre, branch, description, invoice date) from the month's Pastel import file by
invoice number, and writes one JSON file for build_billing_seed.py.

The sheets are free-form: totals add or subtract hand-set cents, some invoices bill a share of a site
("less 30%"), some show a site's full cost for information and total only part of it. So nothing is
re-derived: every row keeps the amount the sheet DISPLAYS (column L), and in_total says whether the
sheet's own "Total Excluding VAT" formula adds that row; the formula's constant cents become the
schedule's total_adjust. Sum(in_total amounts) + total_adjust reproduces finance's figure exactly.

Usage:  python scripts/parse_billing_workbooks.py <master.xlsx> <harmony.xlsx> <pastel.csv> <out.json>
The workbooks hold client billing data: keep the output OUT of the repository (backups/ is ignored).
"""
import csv, json, re, sys
import openpyxl

TEMPLATE_SHEETS = {'1 Page Invoice', '2 Page Invoice', '3 Page Invoice', 'Increase', 'Site',
                   'Input Sheet', 'List', 'Increase March 2024', 'Proforma'}


def txt(v):
    return '' if v is None else str(v).strip()


def num(v):
    if isinstance(v, bool):
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def read_pastel(path):
    out, cur = {}, None
    for r in csv.reader(open(path, newline='', encoding='latin-1')):
        if not r:
            continue
        if r[0] == 'Header':
            cur = {'doc': r[1].strip(), 'customer': r[4].strip(), 'period': r[5].strip(), 'date': r[6].strip(),
                   'order': r[7].strip(), 'branch': r[18].strip(), 'lines': []}
            out[cur['doc']] = cur
        elif r[0] == 'Detail' and cur:
            cur['lines'].append({'excl': float(r[3]), 'incl': float(r[4]), 'tax': r[6].strip(), 'gl': r[9].strip(),
                                 'desc': r[10].strip(), 'cost_centre': r[12].strip()})
    return out


def total_rows(formula):
    """Rows (column L) a "Total Excluding VAT" formula adds, plus its constant cents.
    Handles =SUM(L19:M46,L48:M85), =SUM(..)+SUM(..), =L34+L64, and +/- constants."""
    if not isinstance(formula, str) or not formula.startswith('='):
        return None, 0.0
    body = formula[1:].replace(' ', '')
    rows = set()
    for a, b in re.findall(r'L(\d+):M(\d+)', body):
        rows.update(range(int(a), int(b) + 1))
    rest = re.sub(r'SUM\([^)]*\)', '', body)
    for r in re.findall(r'L(\d+)', rest):
        rows.add(int(r))
    const = 0.0
    for sign, n in re.findall(r'([+-])(\d+(?:\.\d+)?)', re.sub(r'L\d+', '', rest)):
        const += float(n) * (1 if sign == '+' else -1)
    return rows, const


def parse_sheet(ws_v, ws_f, source):
    """One invoice sheet -> schedule dict, or None when it is not an invoice."""
    inv = num(ws_v['L4'].value)
    if not inv:
        return None
    hdr = {}
    for row in range(9, 15):
        lab = txt(ws_v.cell(row, 10).value).rstrip(':').lower()
        if lab:
            hdr[lab] = txt(ws_v.cell(row, 12).value)
    addr = [txt(ws_v.cell(r, 3).value) for r in range(9, 14)]
    addr = [a for a in addr if a]
    total_row = next((r for r in range(19, ws_v.max_row + 1)
                      if txt(ws_v.cell(r, 9).value).lower().startswith('total excluding vat')), None)
    if not total_row:
        return None
    counted, const = total_rows(ws_f.cell(total_row, 12).value)
    lines = []
    for row in range(19, total_row):
        b = ws_v.cell(row, 2).value
        if txt(b).lower().startswith('payment details'):
            break
        desc = txt(ws_v.cell(row, 5).value)
        l_v, l_f = ws_v.cell(row, 12).value, ws_f.cell(row, 12).value
        if txt(l_v).lower().startswith('page '):
            continue
        amt = num(l_v)
        qty, unit = num(ws_v.cell(row, 4).value), num(ws_v.cell(row, 10).value)
        uf = ws_f.cell(row, 10).value
        base = {'row': row, 'description': desc, 'amount': amt if amt else None,
                'in_total': bool(counted and row in counted and amt)}
        lf = l_f if isinstance(l_f, str) and l_f.startswith('=') else None
        if lf and lf.upper().startswith('=SUM'):
            if amt or desc:
                lines.append(dict(base, kind='subtotal', amount_formula=lf))
        elif num(b) is not None or (amt and (qty is not None or unit is not None)):
            lines.append(dict(base, kind='item', item_no=int(num(b)) if num(b) is not None else None,
                              qty=qty, unit_price=unit,
                              unit_formula=uf if isinstance(uf, str) and uf.startswith('=') else None,
                              amount_formula=lf))
        elif amt:
            lines.append(dict(base, kind='amount', amount_formula=lf))
        elif desc:
            font = ws_v.cell(row, 5).font
            lines.append(dict(base, kind='heading' if (font and (font.underline or font.italic)) else 'text'))
    counted_total = sum(l['amount'] or 0 for l in lines if l['in_total'])
    return {
        'source_file': source, 'sheet': ws_v.title, 'invoice_no': int(inv),
        'invoice_date_text': txt(ws_v['J7'].value), 'title': txt(ws_v['B16'].value), 'month_text': txt(ws_v['B17'].value),
        'bill_to_name': addr[0] if addr else '', 'bill_to_address': '\n'.join(addr[1:]),
        'client_vat_no': txt(ws_v['G14'].value), 'order_no': hdr.get('order no', ''),
        'ref_label': next((k.title() for k in hdr if k in ('proposal no', 'contract no')), 'Proposal No'),
        'ref_value': hdr.get('proposal no', hdr.get('contract no', '')),
        'project_no': hdr.get('project no', ''), 'customer_code': hdr.get('customer no', ''),
        'lines': lines, 'total_adjust': round(const, 6), 'total_formula': ws_f.cell(total_row, 12).value,
        'items_total': round(counted_total + const, 6), 'sheet_total': num(ws_v.cell(total_row, 12).value),
    }


def main(master, harmony, pastel_csv, out):
    pastel = read_pastel(pastel_csv)
    schedules, skipped = [], []
    for path in (master, harmony):
        wv = openpyxl.load_workbook(path, data_only=True)
        wf = openpyxl.load_workbook(path, data_only=False)
        site_notes = {}
        if 'Site' in wv.sheetnames:   # master file index: escalation % and notes per invoice
            for r in wv['Site'].iter_rows(min_row=2, values_only=True):
                m = re.match(r'INV(\d+)', txt(r[0]))
                if m:
                    site_notes[int(m.group(1))] = {'index_name': txt(r[0]), 'psira_pct': num(r[3]), 'cpi_pct': num(r[4]),
                                                   'escalation_month': txt(r[6]), 'note': txt(r[7])}
        if 'List' in wv.sheetnames:   # Harmony index: client cost code per invoice
            for r in wv['List'].iter_rows(min_row=4, values_only=True):
                if len(r) > 13 and num(r[13]):
                    site_notes[int(num(r[13]))] = {'index_name': txt(r[12]), 'client_cost_code': txt(r[4]),
                                                   'business_unit': txt(r[2]), 'site_name': txt(r[3])}
        for name in wv.sheetnames:
            if name in TEMPLATE_SHEETS:
                continue
            s = parse_sheet(wv[name], wf[name], path.replace('\\', '/').split('/')[-1])
            if not s:
                skipped.append(name)
                continue
            s.update(site_notes.get(s['invoice_no'], {}))
            p = pastel.get('INV%d' % s['invoice_no'])
            if p:
                main_line = next((l for l in p['lines'] if l['desc'] != 'VAT adjustment'), p['lines'][0])
                s['pastel'] = {'customer': p['customer'], 'branch': p['branch'], 'date': p['date'], 'period': p['period'],
                               'gl': main_line['gl'], 'cost_centre': main_line['cost_centre'], 'desc': main_line['desc'],
                               'excl': round(sum(l['excl'] for l in p['lines'] if l['desc'] != 'VAT adjustment'), 2)}
            schedules.append(s)
    json.dump({'schedules': schedules, 'skipped_sheets': skipped}, open(out, 'w', encoding='utf8'), indent=1, default=str)
    bad = [s for s in schedules if s['sheet_total'] is not None and abs(s['items_total'] - s['sheet_total']) > 0.005]
    nop = [s for s in schedules if 'pastel' not in s]
    mism = [s for s in schedules if 'pastel' in s and abs(round(s['sheet_total'] or 0, 2) - s['pastel']['excl']) > 0.005]
    print('schedules', len(schedules), '| skipped sheets', skipped)
    print('reproduces its own sheet total:', len(schedules) - len(bad), '| not:', len(bad),
          [(s['sheet'], s['items_total'], s['sheet_total'], s['total_formula']) for s in bad][:10])
    print('no Pastel line for the invoice no:', len(nop), [(s['sheet'], s['invoice_no']) for s in nop][:12])
    print('sheet total differs from the Pastel file:', len(mism),
          [(s['sheet'], s['invoice_no'], round(s['sheet_total'], 2), s['pastel']['excl']) for s in mism][:12])


if __name__ == '__main__':
    main(*sys.argv[1:5])
